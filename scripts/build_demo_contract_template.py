"""Build the DOCX asset with python-docx; placeholders match contract_template_fields()."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def build():
    doc = Document()
    for element in (doc.styles.element, doc.element):
        for border in list(element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    doc.core_properties.title = "Базовый шаблон договора СОВА"
    doc.core_properties.author = "СОВА"

    def p(text):
        doc.add_paragraph(text)

    def h(text):
        doc.add_heading(text, level=1)

    doc.add_paragraph("Договор о сотрудничестве и обучении", "Title")
    p("Номер {{ contract_number | default('', true) }}    Дата {{ contract_date | default('', true) }}")
    p("Место заключения: {{ city | default('', true) }}")
    h("1 Стороны договора")
    p("Исполнитель: __________________________________________________________, "
      "в лице ________________________________, действующего на основании __________________.")
    p("Заказчик: {{ counterparty.name | default('', true) }}. "
      "Краткое наименование: {{ counterparty.short_name | default('', true) }}.")
    p("Подписант заказчика: {{ signatory.full_name | default('', true) }}. "
      "Должность: {{ signatory.position | default('', true) }}. "
      "Основание полномочий: {{ signatory.basis | default('', true) }}.")
    h("2 Предмет и состав работ")
    p("Исполнитель и Заказчик согласуют обучение и связанные работы в объёме, указанном "
      "в спецификации на следующей странице. Состав программ, продуктов и лицензий определяется "
      "выбранными сторонами позициями спецификации.")
    p("Формат и место обучения: __________________________________________________.\n"
      "Количество участников и объём часов: _______________________________________.")
    h("3 Сроки и порядок исполнения")
    p("Начало: __________________. Окончание: __________________.\n"
      "Расписание, порядок доступа и ответственные: _______________________________.")
    p("Исполнитель организует согласованные работы. Заказчик предоставляет необходимые "
      "сведения и согласует результаты. Изменения объёма и сроков стороны оформляют письменно.")
    h("4 Стоимость и расчёты")
    p("Сумма договора: {{ amount | default('', true) }}. Валюта: __________________.\n"
      "НДС и порядок его учёта: __________________________________________________.\n"
      "Сроки и порядок оплаты: ___________________________________________________.")
    h("5 Приёмка и дополнительные условия")
    p("Документ и срок приёмки: __________________________________________________.\n"
      "Ответственность сторон и порядок урегулирования разногласий: __________________\n"
      "_________________________________________________________________________.")
    p("Дополнительные условия: {{ comment | default('', true) }}")

    doc.add_page_break()
    doc.add_paragraph("Спецификация и реквизиты", "Title")
    p("К договору № {{ contract_number | default('', true) }} от {{ contract_date | default('', true) }}")
    h("6 Направления и программы")
    p("Направления: {% for direction in directions %}{{ direction.name }}"
      "{% if not loop.last %}; {% endif %}{% else %}не выбраны{% endfor %}.")
    p("{% for program in programs %}{{ program.name }} — {{ program.direction | default('', true) }}"
      "{% if not loop.last %}; {% endif %}{% else %}Программы не выбраны{% endfor %}.")
    h("7 Продукты и лицензии")
    p("{% for product in products %}{{ product.name }}"
      "{% if product.program | default('') %} (программа: {{ product.program }}){% endif %}"
      "{% if not loop.last %}; {% endif %}{% else %}Продукты не выбраны{% endfor %}.")
    p("{% for license in licenses %}Продукт: {{ license.product }}; "
      "договор лицензии: {{ license.contract_number | default('', true) }}; "
      "дата подписания: {{ license.signed_at | default('', true) }}; "
      "действует до года: {{ license.valid_until_year | default('', true) }}; "
      "подписана: {{ 'да' if license.is_signed else 'нет' }}. "
      "{% else %}Лицензии не выбраны.{% endfor %}")
    h("8 Реквизиты исполнителя")
    p("Наименование: __________________________________________________________.\n"
      "ИНН / КПП: ________________________. Адрес: ______________________________.\n"
      "Банк / БИК: ____________________________________________________________.\n"
      "Расчётный / корреспондентский счёт: ______________________________________.\n"
      "Email / телефон: _______________________________________________________.")
    h("9 Реквизиты заказчика")
    p("Наименование или ФИО: {{ counterparty.name | default('', true) }}\n"
      "ИНН: {{ counterparty.inn | default('', true) }}\n"
      "Адрес: {{ counterparty.address | default('', true) }}\n"
      "Email: {{ counterparty.email | default('', true) }}\n"
      "Телефон: {{ counterparty.phone | default('', true) }}")
    p("Банковские реквизиты при необходимости: ___________________________________.")
    h("10 Подписи сторон")
    p("Исполнитель: __________________ / __________________________ /\n"
      "Заказчик: __________________ / {{ signatory.full_name | default('', true) }} /")
    path = Path(__file__).resolve().parents[1] / "reference_data" / "contract_base.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    build()
