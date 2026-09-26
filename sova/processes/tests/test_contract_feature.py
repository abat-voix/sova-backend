from io import BytesIO

from django.core.files.base import ContentFile
from docx import Document
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import SystemRole, UserRole

from sova.catalog.tests.factories import ContactPersonFactory, UniversityFactory
from sova.core.tests.factories import UserFactory
from sova.core.tests.media import TemporaryMediaMixin
from sova.interactions.enum import DocumentTemplateKind
from sova.interactions.models import Contract, DocumentTemplate, InteractionContact
from sova.interactions.tests.factories import (
    ContractFactory,
    InteractionDirectionFactory,
    InteractionFactory,
    InteractionProductFactory,
    InteractionProgramFactory,
    LicenseFactory,
)
from sova.processes.models import ActionFeatureExecution
from sova.processes.tests.factories import ActionInstanceFactory
from sova.workflows.tests.factories import ActionFeatureFactory


class CreateContractFeatureApiTestCase(TemporaryMediaMixin, APITestCase):
    """Feature `contract.create`: выбор шаблона и генерация договора из JSON."""

    def setUp(self) -> None:
        self.user = UserFactory()
        UserRole.objects.create(user=self.user, role=SystemRole.PLATFORM_ADMIN)
        self.client.force_authenticate(user=self.user)

        university = UniversityFactory(short_name="ВУЗ", inn="7700000000", city="Москва")
        self.interaction = InteractionFactory(university=university)
        self.action_instance = ActionInstanceFactory(
            stage_instance__workflow_instance__interaction=self.interaction,
        )
        ActionFeatureFactory(action=self.action_instance.action, code="contract.create")
        self.template = DocumentTemplate.objects.get(kind=DocumentTemplateKind.CONTRACT)
        self._set_template(
            "№ {{ contract_number }} от {{ contract_date }}",
            "{{ counterparty.name }} / {{ signatory.full_name }}",
            "{% for product in products %}{{ product.name }} {% endfor %}",
            "Сумма: {{ amount }}",
        )
        self.base_url = f"/api/processes/action-instances/{self.action_instance.pk}/features/contract.create"

    def _set_template(self, *paragraphs: str) -> None:
        source = Document()
        for paragraph in paragraphs:
            source.add_paragraph(paragraph)
        content = BytesIO()
        source.save(content)
        self.template.file.save("contract.docx", ContentFile(content.getvalue()), save=True)

    def payload(self, **document) -> dict:
        return {
            "template": str(self.template.pk),
            "document": {
                "contract_number": "Д-1",
                "contract_date": "2026-09-26",
                "counterparty": {"name": "Вуз", "inn": "7700000000"},
                "signatory": {"full_name": "Иванов И. И.", "position": "Ректор"},
                "amount": "1000.50",
                **document,
            },
        }

    def test_initial_prefills_counterparty_contacts_and_products(self) -> None:
        contact = ContactPersonFactory(university=self.interaction.university, position="Ректор")
        InteractionContact.objects.create(interaction=self.interaction, contact_person=contact)
        product = InteractionProductFactory(interaction=self.interaction)

        response = self.client.get(f"{self.base_url}/initial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        self.assertEqual(response.data["templates"], [{"id": self.template.pk, "name": "Договор"}])
        self.assertEqual(response.data["contacts"][0]["full_name"], contact.full_name)
        document = response.data["document"]
        self.assertEqual(document["counterparty"]["name"], self.interaction.university.name)
        self.assertEqual(document["counterparty"]["short_name"], "ВУЗ")
        self.assertEqual(document["counterparty"]["inn"], "7700000000")
        self.assertEqual(document["city"], "Москва")
        self.assertEqual(
            document["products"], [{"id": str(product.pk), "name": product.product.name, "program": ""}],
        )

    def test_initial_contains_directions_programs_and_licenses(self) -> None:
        direction = InteractionDirectionFactory(interaction=self.interaction)
        InteractionDirectionFactory(interaction=self.interaction, is_active=False)
        InteractionDirectionFactory()  # чужое взаимодействие
        program = InteractionProgramFactory(interaction=self.interaction)
        product = InteractionProductFactory(interaction=self.interaction, interaction_program=program)
        license_ = LicenseFactory(
            contract=ContractFactory(interaction=self.interaction, contract_number="Л-1"),
            interaction_product=product,
            valid_until_year=2027,
        )

        response = self.client.get(f"{self.base_url}/initial/")

        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        document = response.data["document"]
        self.assertEqual(document["directions"], [{"id": str(direction.pk), "name": direction.direction.name}])
        self.assertEqual(
            document["programs"],
            [{"id": str(program.pk), "name": program.program.name, "direction": program.program.direction.name}],
        )
        self.assertEqual(document["products"][0]["program"], program.program.name)
        self.assertEqual(
            document["licenses"],
            [{
                "id": str(license_.pk),
                "product": product.product.name,
                "contract_number": "Л-1",
                "signed_at": None,
                "valid_until_year": 2027,
                "is_signed": False,
            }],
        )

    def test_execute_resolves_selected_scope_from_interaction(self) -> None:
        direction = InteractionDirectionFactory(interaction=self.interaction)
        InteractionDirectionFactory(interaction=self.interaction)
        product = InteractionProductFactory(interaction=self.interaction)

        response = self.client.post(
            f"{self.base_url}/execute/",
            self.payload(
                directions=[{"id": str(direction.pk), "name": "подменённое имя"}],
                products=[{"id": str(product.pk)}],
            ),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        document = response.data["target"]["data"]["document"]
        self.assertEqual(document["directions"], [{"id": str(direction.pk), "name": direction.direction.name}])
        self.assertEqual([item["id"] for item in document["products"]], [str(product.pk)])
        self.assertEqual(document["programs"], [])
        self.assertEqual(document["licenses"], [])

    def test_execute_rejects_scope_of_another_interaction(self) -> None:
        foreign = InteractionProductFactory()

        response = self.client.post(
            f"{self.base_url}/execute/", self.payload(products=[{"id": str(foreign.pk)}]), format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("products", response.data["document"])
        self.assertFalse(Contract.objects.exists())

    def test_execute_renders_json_and_creates_contract_file(self) -> None:
        response = self.client.post(f"{self.base_url}/execute/", self.payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        contract = Contract.objects.get(pk=response.data["target"]["id"])
        self.assertEqual(contract.interaction, self.interaction)
        self.assertEqual(contract.contract_number, "Д-1")
        self.assertTrue(contract.file)
        self.assertEqual(contract.file_name, "Договор.docx")
        with contract.file.open("rb") as generated:
            text = "\n".join(paragraph.text for paragraph in Document(generated).paragraphs)
        self.assertIn("№ Д-1 от 2026-09-26", text)
        self.assertIn("Вуз / Иванов И. И.", text)
        self.assertIn("Сумма: 1000.50", text)
        self.assertNotIn("{{", text)
        self.assertEqual(contract.files.count(), 1)
        self.assertEqual(contract.files.get().uploaded_by, self.user)
        data = response.data["target"]["data"]
        self.assertTrue(data["file_generated"])
        self.assertEqual(data["document"]["contract_date"], "2026-09-26")
        self.assertEqual(data["document"]["amount"], "1000.50")
        self.assertEqual(data["document"]["signatory"]["full_name"], "Иванов И. И.")
        execution = ActionFeatureExecution.objects.get(pk=response.data["execution"]["id"])
        self.assertEqual(execution.result["document"]["counterparty"]["name"], "Вуз")

    def test_execute_renders_selected_product_and_escapes_markup(self) -> None:
        product = InteractionProductFactory(interaction=self.interaction)
        response = self.client.post(
            f"{self.base_url}/execute/",
            self.payload(counterparty={"name": "ООО <Тест>"}, products=[{"id": str(product.pk)}]),
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        contract = Contract.objects.get(pk=response.data["target"]["id"])
        with contract.file.open("rb") as generated:
            text = "\n".join(paragraph.text for paragraph in Document(generated).paragraphs)
        self.assertIn("ООО <Тест>", text)
        self.assertIn(product.product.name, text)

    def test_execute_renders_empty_optional_value_as_blank(self) -> None:
        response = self.client.post(
            f"{self.base_url}/execute/", self.payload(amount=None), format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        contract = Contract.objects.get(pk=response.data["target"]["id"])
        with contract.file.open("rb") as generated:
            text = "\n".join(paragraph.text for paragraph in Document(generated).paragraphs)
        self.assertIn("Сумма: ", text)
        self.assertNotIn("None", text)

    def test_execute_rejects_template_without_file(self) -> None:
        self.template.file.delete(save=True)
        initial = self.client.get(f"{self.base_url}/initial/")
        self.assertEqual(initial.data["templates"], [])
        response = self.client.post(f"{self.base_url}/execute/", self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("template", response.data)
        self.assertFalse(Contract.objects.exists())

    def test_execute_rejects_corrupt_docx(self) -> None:
        self.template.file.save("broken.docx", ContentFile(b"not a docx"), save=True)
        response = self.client.post(f"{self.base_url}/execute/", self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("template", response.data)
        self.assertFalse(Contract.objects.exists())

    def test_execute_rejects_unfilled_template_variable(self) -> None:
        self._set_template("{{ nonexistent_field }}")
        response = self.client.post(f"{self.base_url}/execute/", self.payload(), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("template", response.data)
        self.assertFalse(Contract.objects.exists())

    def test_execute_requires_counterparty_name(self) -> None:
        response = self.client.post(
            f"{self.base_url}/execute/", self.payload(counterparty={"name": ""}), format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(Contract.objects.exists())

    def test_execute_rejects_inactive_template(self) -> None:
        self.template.is_active = False
        self.template.save()

        response = self.client.post(f"{self.base_url}/execute/", self.payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(Contract.objects.exists())

    def test_available_features_include_contract_create(self) -> None:
        response = self.client.get(f"/api/processes/action-instances/{self.action_instance.pk}/")

        self.assertEqual(response.status_code, status.HTTP_200_OK, msg=response.data)
        self.assertEqual([item["code"] for item in response.data["available_features"]], ["contract.create"])
