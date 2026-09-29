import io
import json
import shutil
import tempfile
from unittest import mock

import xlrd
from django.core.files.storage import FileSystemStorage
from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import SystemRole, UserRole
from sova.catalog.tests.factories import ProductFactory, ProgramFactory, OrganizationFactory
from sova.core.tests.factories import UserFactory
from sova.interactions.models import Responsible
from sova.interactions.services import responsible_service
from sova.interactions.tests.factories import (
    InteractionDirectionFactory,
    InteractionFactory,
    InteractionProductFactory,
    InteractionProgramFactory,
)
from sova.processes.tests.factories import StageInstanceFactory, WorkflowInstanceFactory
from sova.reports.enum import ReportJobStatus
from sova.reports.models import ReportJob
from sova.reports.services import jobs


class ReportTestMixin:
    """Общие данные: администратор и типовое взаимодействие с составом."""

    preview_url = reverse("reports:interaction-report-preview")
    summary_url = reverse("reports:interaction-report-summary")
    export_url = reverse("reports:interaction-report-export")

    @staticmethod
    def create_user(role=None):
        user = UserFactory()
        if role is not None:
            UserRole.objects.create(user=user, role=role)
        return user

    def setUp(self) -> None:
        self.admin = self.create_user(SystemRole.PLATFORM_ADMIN)
        self.client.force_authenticate(user=self.admin)

    def preview(self, **data):
        response = self.client.post(self.preview_url, data={"page_size": 200, **data}, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        return response.data


class InteractionReportRowsTestCase(ReportTestMixin, APITestCase):
    """Строки предпросмотра: единица детализации и фильтры состава."""

    def test_empty_interaction_gives_one_row(self) -> None:
        interaction = InteractionFactory()

        data = self.preview()

        self.assertEqual(data["count"], 1)
        row = data["results"][0]
        self.assertEqual(row["interaction_id"], str(interaction.pk))
        self.assertIsNone(row["product_id"])
        self.assertEqual(row["active_stages"], [])
        self.assertEqual(row["process_status"], [])

    def test_products_programs_and_directions(self) -> None:
        """Продукт с программой, продукт без программы, программа без продуктов, направление без программ."""
        interaction = InteractionFactory()
        with_products = InteractionProgramFactory(interaction=interaction)
        without_products = InteractionProgramFactory(interaction=interaction)
        product = ProductFactory(programs=[with_products.program])
        InteractionProductFactory(interaction=interaction, product=product, interaction_program=with_products)
        loose = InteractionProductFactory(interaction=interaction)
        # Направление программы уже покрыто строкой программы — отдельной строки не даёт
        InteractionDirectionFactory(interaction=interaction, direction=with_products.program.direction)
        lonely_direction = InteractionDirectionFactory(interaction=interaction)

        rows = self.preview()["results"]

        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0]["product_id"], str(product.pk))
        self.assertEqual(rows[0]["program_id"], str(with_products.program_id))
        self.assertEqual(rows[0]["direction_id"], str(with_products.program.direction_id))
        self.assertIsNotNone(rows[0]["interaction_direction_id"])
        self.assertEqual(rows[1]["product_id"], str(loose.product_id))
        self.assertIsNone(rows[1]["program_id"])
        self.assertEqual(rows[1]["direction"], "")
        self.assertEqual(rows[2]["program_id"], str(without_products.program_id))
        self.assertIsNone(rows[2]["product_id"])
        self.assertEqual(rows[3]["direction_id"], str(lonely_direction.direction_id))

    def test_inactive_composition_is_skipped(self) -> None:
        interaction = InteractionFactory()
        InteractionProductFactory(interaction=interaction, is_active=False)

        rows = self.preview()["results"]

        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]["product_id"])

    def test_product_filter_limits_rows(self) -> None:
        """Выбранное взаимодействие не приносит в отчёт другие продукты."""
        interaction = InteractionFactory()
        wanted = InteractionProductFactory(interaction=interaction)
        InteractionProductFactory(interaction=interaction)
        InteractionFactory()

        data = self.preview(products=[str(wanted.product_id)])

        self.assertEqual(data["count"], 1)
        self.assertEqual(data["results"][0]["interaction_product_id"], str(wanted.pk))

    def test_direction_and_program_filters_intersect(self) -> None:
        interaction = InteractionFactory()
        program = InteractionProgramFactory(interaction=interaction)
        InteractionProgramFactory(interaction=interaction)
        other_program = ProgramFactory()

        data = self.preview(
            directions=[str(program.program.direction_id)], programs=[str(program.program_id)]
        )
        self.assertEqual([row["interaction_program_id"] for row in data["results"]], [str(program.pk)])

        data = self.preview(
            directions=[str(program.program.direction_id)], programs=[str(other_program.pk)]
        )
        self.assertEqual(data["count"], 0)

    def test_parallel_active_stages_and_process_status(self) -> None:
        interaction = InteractionFactory()
        product = InteractionProductFactory(interaction=interaction)
        process = WorkflowInstanceFactory(interaction=interaction)
        StageInstanceFactory(workflow_instance=process, context_type="interaction", context_id=None)
        StageInstanceFactory(workflow_instance=process, context_type="product", context_id=product.pk)
        StageInstanceFactory(
            workflow_instance=process, context_type="product", context_id=product.pk, status="completed"
        )

        data = self.preview()

        row = data["results"][0]
        self.assertEqual(len(row["active_stages"]), 2)
        self.assertEqual(row["process_status"][0]["status"], "running")
        self.assertEqual(data["meta"]["active_stage_status"], "in_progress")

    def test_current_responsible_only(self) -> None:
        interaction = InteractionFactory()
        old, new = self.create_user(SystemRole.KAM), self.create_user(SystemRole.KAM)
        responsible_service.assign(interaction=interaction, manager=old, assigned_by=None)
        Responsible.objects.filter(interaction=interaction).update(unassigned_at="2026-01-01T00:00:00Z")
        responsible_service.assign(interaction=interaction, manager=new, assigned_by=None)

        self.assertEqual(self.preview()["results"][0]["responsible_ids"], [new.pk])
        self.assertEqual(self.preview(responsibles=[old.pk])["count"], 0)

    def test_several_responsibles(self) -> None:
        """Все действующие КАМы в строке по алфавиту; фильтр находит взаимодействие по любому из них."""
        interaction = InteractionFactory()
        yakovlev = UserFactory(first_name="Яков", last_name="Яковлев")
        antonov = UserFactory(first_name="Антон", last_name="Антонов")
        for manager in (yakovlev, antonov):
            responsible_service.assign(interaction=interaction, manager=manager, assigned_by=None)

        row = self.preview()["results"][0]

        self.assertEqual(row["responsible_ids"], [antonov.pk, yakovlev.pk])
        self.assertEqual(row["responsible"], ["Антон Антонов", "Яков Яковлев"])
        self.assertEqual(self.preview(responsibles=[yakovlev.pk])["count"], 1)

    def test_ordering_by_first_responsible(self) -> None:
        """Сортировка по ответственному — по первому КАМу по алфавиту, без КАМа — в конце."""
        several = InteractionFactory()
        for last_name in ("Яковлев", "Антонов"):
            responsible_service.assign(
                interaction=several, manager=UserFactory(last_name=last_name), assigned_by=None
            )
        single = InteractionFactory()
        responsible_service.assign(interaction=single, manager=UserFactory(last_name="Борисов"), assigned_by=None)
        nobody = InteractionFactory()

        rows = self.preview(ordering="responsible")["results"]

        self.assertEqual(
            [row["interaction_id"] for row in rows], [str(several.pk), str(single.pk), str(nobody.pk)]
        )

    def test_selected_columns_only(self) -> None:
        InteractionFactory()

        data = self.preview(columns=["organization"])

        row = data["results"][0]
        self.assertIn("organization", row)
        self.assertIn("interaction_id", row)
        self.assertNotIn("product", row)
        self.assertEqual([c["key"] for c in data["meta"]["columns"]], ["organization"])

    def test_period_by_created_at(self) -> None:
        old = InteractionFactory()
        type(old).objects.filter(pk=old.pk).update(created_at="2025-01-10T12:00:00Z")
        recent = InteractionFactory()
        type(recent).objects.filter(pk=recent.pk).update(created_at="2025-02-10T12:00:00Z")

        data = self.preview(date_from="2025-02-01", date_to="2025-02-10")

        self.assertEqual([row["interaction_id"] for row in data["results"]], [str(recent.pk)])
        self.assertEqual(data["meta"]["period_basis"], "interaction_created_at")

    def test_pagination(self) -> None:
        for _ in range(3):
            InteractionFactory()

        response = self.client.post(self.preview_url, data={"page": 2, "page_size": 2}, format="json")

        self.assertEqual(response.data["count"], 3)
        self.assertEqual(len(response.data["results"]), 1)


class InteractionReportValidationTestCase(ReportTestMixin, APITestCase):
    """Неверные фильтры дают 400 с кодами ошибок."""

    def test_null_columns_uses_all_columns(self) -> None:
        for url in (self.preview_url, self.summary_url):
            with self.subTest(url=url):
                response = self.client.post(url, data={"columns": None}, format="json")
                self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
                self.assertEqual(response.data["meta"]["columns"][0]["key"], "organization")

    def assert_error(self, data, field, code):
        response = self.client.post(self.preview_url, data=data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data[field][0].code, code)

    def test_period_order(self) -> None:
        self.assert_error({"date_from": "2025-02-01", "date_to": "2025-01-01"}, "date_to", "period_order")

    def test_unknown_ids(self) -> None:
        self.assert_error(
            {"organizations": ["00000000-0000-0000-0000-000000000000"]}, "organizations", "unknown_ids"
        )

    def test_unknown_column(self) -> None:
        response = self.client.post(self.preview_url, data={"columns": ["password"]}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_page_size_limit(self) -> None:
        response = self.client.post(self.preview_url, data={"page_size": 1000}, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class InteractionReportVisibilityTestCase(ReportTestMixin, APITestCase):
    """Роли: предпросмотр и статистика не раскрывают чужие взаимодействия."""

    def test_kam_sees_only_own_rows_and_counts(self) -> None:
        kam = self.create_user(SystemRole.KAM)
        own = InteractionFactory()
        responsible_service.assign(interaction=own, manager=kam, assigned_by=None)
        foreign = InteractionFactory()
        responsible_service.assign(interaction=foreign, manager=self.create_user(SystemRole.KAM), assigned_by=None)
        self.client.force_authenticate(user=kam)

        rows = self.preview()["results"]
        summary = self.client.post(self.summary_url, data={}, format="json").data

        self.assertEqual([row["interaction_id"] for row in rows], [str(own.pk)])
        self.assertEqual(summary["interactions_count"], 1)
        self.assertNotIn(str(foreign.organization_id), [item["id"] for item in summary["by_organization"]])

    def test_user_without_role_gets_nothing(self) -> None:
        InteractionFactory()
        self.client.force_authenticate(user=self.create_user())

        response = self.client.post(self.preview_url, data={"page_size": 200}, format="json")
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_anonymous_rejected(self) -> None:
        self.client.force_authenticate(user=None)
        response = self.client.post(self.preview_url, data={}, format="json")
        self.assertIn(response.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))


class InteractionReportSummaryTestCase(ReportTestMixin, APITestCase):
    def test_unique_interaction_counts(self) -> None:
        """Вуз с несколькими продуктами не завышает число взаимодействий."""
        organization = OrganizationFactory()
        interaction = InteractionFactory(organization=organization)
        InteractionProductFactory(interaction=interaction)
        InteractionProductFactory(interaction=interaction)
        InteractionFactory(organization=organization)

        data = self.client.post(self.summary_url, data={}, format="json").data

        self.assertEqual(data["interactions_count"], 2)
        self.assertEqual(data["rows_count"], 3)
        self.assertEqual(data["products_count"], 2)
        self.assertEqual(data["by_organization"][0]["interactions"], 2)
        self.assertEqual(data["by_process_status"][0]["status"], None)

    def test_interaction_counts_for_each_responsible(self) -> None:
        """Взаимодействие с двумя КАМами попадает в группу каждого; без КАМа — в «Без ответственного»."""
        first, second = UserFactory(), UserFactory()
        shared = InteractionFactory()
        for manager in (first, second):
            responsible_service.assign(interaction=shared, manager=manager, assigned_by=None)
        InteractionFactory()

        data = self.client.post(self.summary_url, data={}, format="json").data

        groups = {item["id"]: item["interactions"] for item in data["by_responsible"]}
        self.assertEqual(groups, {str(first.pk): 1, str(second.pk): 1, None: 1})
        self.assertEqual(data["interactions_count"], 2)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True)
class ReportExportTestCase(ReportTestMixin, APITestCase):
    """Выгрузки: одинаковые строки во всех форматах, доступ только владельцу."""

    def setUp(self) -> None:
        super().setUp()
        self.storage_dir = tempfile.mkdtemp()
        field = ReportJob._meta.get_field("file")
        self.storage_patch = mock.patch.object(field, "storage", FileSystemStorage(location=self.storage_dir))
        self.storage_patch.start()

    def tearDown(self) -> None:
        self.storage_patch.stop()
        shutil.rmtree(self.storage_dir, ignore_errors=True)

    def export(self, report_format, **data):
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(self.export_url, data={"format": report_format, **data}, format="json")
        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED, response.data)
        job = ReportJob.objects.get(pk=response.data["id"])
        self.assertEqual(job.status, ReportJobStatus.READY, job.error_message)
        return job

    def download(self, job) -> bytes:
        response = self.client.get(reverse("reports:report-job-download", args=[job.pk]))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return b"".join(response.streaming_content)

    def setup_data(self):
        organization = OrganizationFactory(name="Московский университет")
        interaction = InteractionFactory(organization=organization)
        InteractionProductFactory(interaction=interaction)
        InteractionProductFactory(interaction=interaction)
        InteractionFactory()

    def test_formats_have_same_rows(self) -> None:
        self.setup_data()
        preview = self.preview(columns=["organization", "product"])
        expected = [[row["organization"], row["product"]] for row in preview["results"]]

        content = self.download(self.export("xlsx", columns=["organization", "product"]))
        sheet = load_workbook(io.BytesIO(content)).active
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        header_index = values.index(["Организация", "Продукт"])
        self.assertEqual([[v or "" for v in row] for row in values[header_index + 1:]], expected)

        content = self.download(self.export("xls", columns=["organization", "product"]))
        sheet = xlrd.open_workbook(file_contents=content).sheet_by_index(0)
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        header_index = rows.index(["Организация", "Продукт"])
        self.assertEqual(rows[header_index + 1:], expected)

        content = self.download(self.export("json", columns=["organization", "product"]))
        data = json.loads(content)
        self.assertEqual([[r["organization"], r["product"]] for r in data["rows"]], expected)
        self.assertIn("interaction_product_id", data["rows"][0])

    def test_pdf_uses_gotenberg(self) -> None:
        self.setup_data()
        with mock.patch("sova.reports.services.exporters.html_to_pdf", return_value=b"%PDF-1.7") as convert:
            job = self.export("pdf")
        html = convert.call_args.args[0]
        self.assertIn("Московский университет", html)
        self.assertEqual(self.download(job), b"%PDF-1.7")

    def test_foreign_job_is_not_found(self) -> None:
        job = self.export("json")
        self.client.force_authenticate(user=self.create_user(SystemRole.PLATFORM_ADMIN))

        detail = self.client.get(reverse("reports:report-job-detail", args=[job.pk]))
        download = self.client.get(reverse("reports:report-job-download", args=[job.pk]))

        self.assertEqual(detail.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(download.status_code, status.HTTP_404_NOT_FOUND)

    def test_job_rechecks_visibility(self) -> None:
        """Worker строит отчёт с текущей видимостью владельца."""
        kam = self.create_user(SystemRole.KAM)
        foreign = InteractionFactory()
        responsible_service.assign(interaction=foreign, manager=self.create_user(SystemRole.KAM), assigned_by=None)
        self.client.force_authenticate(user=kam)

        job = self.export("json", organizations=[str(foreign.organization_id)])

        self.assertEqual(json.loads(self.download(job))["rows"], [])

    def test_not_ready_download_conflict(self) -> None:
        job = ReportJob.objects.create(owner=self.admin, spec={}, format="xlsx")
        response = self.client.get(reverse("reports:report-job-download", args=[job.pk]))
        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(response.data["code"], "not_ready")

    def test_rerun_is_noop_and_failure_is_safe(self) -> None:
        job = self.export("json")
        jobs.run_job(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.attempts, 1)

        with mock.patch("sova.reports.services.jobs.export", side_effect=RuntimeError("секрет БД")):
            failed = self.export_failed()
        self.assertEqual(failed.error_code, "internal_error")
        self.assertNotIn("секрет", failed.error_message)

    def export_failed(self):
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(self.export_url, data={"format": "json"}, format="json")
        return ReportJob.objects.get(pk=response.data["id"])

    def test_cleanup_removes_expired(self) -> None:
        job = self.export("json")
        name = job.file.name
        ReportJob.objects.filter(pk=job.pk).update(expires_at="2000-01-01T00:00:00Z")

        jobs.cleanup_jobs()

        self.assertFalse(ReportJob.objects.filter(pk=job.pk).exists())
        self.assertFalse(job.file.storage.exists(name))
