from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

import xlrd
import xlwt
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook

from sova.catalog.exceptions import CatalogImportError
from sova.catalog.enum import CatalogType
from sova.catalog.models import CatalogImportMapping, Direction, Vendor
from sova.catalog.services import catalog_import_service, import_file_service
from sova.catalog.tests.factories import ProductFactory, OrganizationFactory, VendorFactory
from sova.interactions.models import Contract, License


def _write_workbook(path: Path, headers: tuple, *rows: tuple) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def _write_xls(path: Path, headers: tuple, *rows: tuple) -> None:
    """Старый формат .xls (BIFF) — даты пишутся числом со стилем даты, как их сохраняет Excel."""
    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Лист1")
    date_style = xlwt.easyxf(num_format_str="DD.MM.YYYY")
    for row_index, row in enumerate((headers, *rows)):
        for column_index, value in enumerate(row):
            if isinstance(value, date | datetime):
                worksheet.write(row_index, column_index, value, date_style)
            else:
                worksheet.write(row_index, column_index, value)
    workbook.save(str(path))


def _map(catalog_type: str, columns: dict[str, str]) -> None:
    for source_column, target_field in columns.items():
        CatalogImportMapping.objects.create(
            catalog_type=catalog_type, source_column=source_column, target_field=target_field
        )


class CatalogImportFileTestCase(TestCase):
    """import_file переводит произвольные заголовки файла в канонические ключи по CatalogImportMapping."""

    def test_translates_custom_headers_via_mapping(self) -> None:
        _map(CatalogType.DIRECTION, {"Наименование направления": "name", "Код 1С": "external_code"})
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xlsx"
            _write_workbook(source, ("Наименование направления", "Код 1С"), ("Разработка", "dir-1"))

            created, updated, _ = catalog_import_service.import_file(CatalogType.DIRECTION, source)

            self.assertEqual((created, updated), (1, 0))
            self.assertTrue(Direction.objects.filter(name="Разработка", external_code="dir-1").exists())

    def test_ignores_mapping_of_field_removed_from_schema(self) -> None:
        """Сохранённый маппинг поля, которого уже нет в схеме типа (external_code вендора), не применяется."""
        _map(CatalogType.VENDOR, {"Наименование вендора": "name", "Код 1С": "external_code"})
        with TemporaryDirectory() as directory:
            source = Path(directory) / "vendors.xlsx"
            _write_workbook(source, ("Наименование вендора", "Код 1С"), ("Вендор", "vendor-1"))

            created, _, _ = catalog_import_service.import_file(CatalogType.VENDOR, source)

            # Проверяем, что вендор создан без кода из убранной колонки
            self.assertEqual(created, 1)
            self.assertEqual(Vendor.objects.get(name="Вендор").external_code, None)

    def test_multiline_header_matches_single_line_mapping(self) -> None:
        """Заголовок с переносом строки (Alt+Enter в Excel) совпадает с маппингом, записанным в одну строку."""
        _map(CatalogType.VENDOR, {"Наименование вендора": "name"})
        with TemporaryDirectory() as directory:
            source = Path(directory) / "vendors.xlsx"
            _write_workbook(source, ("Наименование\nвендора",), ("JetBrains",))

            created, _, _ = catalog_import_service.import_file(CatalogType.VENDOR, source)

            # Проверяем, что колонка найдена и вендор создан
            self.assertEqual(created, 1)
            self.assertTrue(Vendor.objects.filter(name="JetBrains").exists())

    def test_no_mapping_configured_raises(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "vendors.xlsx"
            _write_workbook(source, ("name",), ("Вендор",))

            with self.assertRaises(CatalogImportError):
                catalog_import_service.import_file(CatalogType.VENDOR, source)

    def test_missing_required_mapped_field_raises(self) -> None:
        CatalogImportMapping.objects.create(
            catalog_type=CatalogType.DIRECTION, source_column="Наименование направления", target_field="name"
        )
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xlsx"
            _write_workbook(source, ("Наименование направления",), ("Направление",))

            with self.assertRaises(CatalogImportError):
                catalog_import_service.import_file(CatalogType.DIRECTION, source)

    def test_dispatches_contract_registry(self) -> None:
        OrganizationFactory(name="МГУ")
        ProductFactory(name="1С:Предприятие", vendor=VendorFactory(name="1С"))
        columns = {"Вуз": "organization", "Вендор": "vendor", "Продукт": "product", "№ договора": "contract_number"}
        for source_column, target_field in columns.items():
            CatalogImportMapping.objects.create(
                catalog_type=CatalogType.CONTRACT_REGISTRY, source_column=source_column, target_field=target_field
            )
        with TemporaryDirectory() as directory:
            source = Path(directory) / "registry.xlsx"
            _write_workbook(source, tuple(columns), ("МГУ", "1С", "1С:Предприятие", "Д-1"))

            created, updated, _ = catalog_import_service.import_file(CatalogType.CONTRACT_REGISTRY, source)

            self.assertEqual((created, updated), (1, 0))
            self.assertIsNone(Contract.objects.get(contract_number="Д-1").interaction_id)


class CatalogImportSourceFormatTestCase(TestCase):
    """Источник импорта: xlsx и xls, путь на диске или загруженный файл."""

    def test_reads_xls_file(self) -> None:
        _map(CatalogType.DIRECTION, {"Направление": "name", "Код": "external_code"})
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xls"
            _write_xls(source, ("Направление", "Код"), ("JetBrains", "jb"), ("1С", "one-c"))

            created, updated, _ = catalog_import_service.import_file(CatalogType.DIRECTION, source)

        self.assertEqual((created, updated), (2, 0))
        self.assertTrue(Direction.objects.filter(name="1С", external_code="one-c").exists())

    def test_xls_integer_number_is_read_without_fraction(self) -> None:
        _map(CatalogType.DIRECTION, {"Направление": "name", "Код": "external_code"})
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xls"
            _write_xls(source, ("Направление", "Код"), ("JetBrains", 1001))

            catalog_import_service.import_file(CatalogType.DIRECTION, source)

        self.assertEqual(Direction.objects.get(name="JetBrains").external_code, "1001")

    def test_xls_date_cell_is_read_as_date(self) -> None:
        OrganizationFactory(name="МГУ")
        ProductFactory(name="1С:Предприятие", vendor=VendorFactory(name="1С"))
        columns = {
            "Вуз": "organization",
            "Вендор": "vendor",
            "Продукт": "product",
            "№ договора": "contract_number",
            "Лицензия подписана": "license_signed",
            "Срок лицензии": "license_valid_until_year",
        }
        _map(CatalogType.CONTRACT_REGISTRY, columns)
        with TemporaryDirectory() as directory:
            source = Path(directory) / "registry.xls"
            _write_xls(source, tuple(columns), ("МГУ", "1С", "1С:Предприятие", "Д-1", date(2026, 3, 1), 2027))

            catalog_import_service.import_file(CatalogType.CONTRACT_REGISTRY, source)

        license_ = License.objects.get(contract__contract_number="Д-1")
        self.assertEqual(license_.signed_at, date(2026, 3, 1))
        self.assertEqual(license_.valid_until_year, 2027)

    def test_accepts_uploaded_file(self) -> None:
        _map(CatalogType.DIRECTION, {"Направление": "name", "Код": "external_code"})
        workbook = Workbook()
        workbook.active.append(("Направление", "Код"))
        workbook.active.append(("JetBrains", "jb"))
        content = BytesIO()
        workbook.save(content)
        upload = SimpleUploadedFile("directions.xlsx", content.getvalue())

        created, updated, _ = catalog_import_service.import_file(CatalogType.DIRECTION, upload)

        self.assertEqual((created, updated), (1, 0))

    def test_unsupported_extension_raises(self) -> None:
        upload = SimpleUploadedFile("directions.csv", b"name\nJetBrains\n")

        with self.assertRaises(CatalogImportError):
            catalog_import_service.import_file(CatalogType.DIRECTION, upload)

    def test_corrupted_file_raises_domain_error(self) -> None:
        _map(CatalogType.DIRECTION, {"Направление": "name", "Код": "external_code"})
        for name in ("directions.xlsx", "directions.xls"):
            with self.subTest(name=name), self.assertRaises(CatalogImportError):
                catalog_import_service.import_file(CatalogType.DIRECTION, SimpleUploadedFile(name, b"not a spreadsheet"))


def _html_xls(body_rows: list[tuple], charset: str | None, encoding: str) -> bytes:
    """HTML-таблица, какой её сохраняют 1С и веб-системы под расширением .xls."""
    meta = f'<meta http-equiv="Content-Type" content="text/html; charset={charset}">' if charset else ""
    rows = "".join("<tr>" + "".join(f"<td>{value}</td>" for value in row) + "</tr>" for row in body_rows)
    return f"<html><head>{meta}</head><body><table>{rows}</table></body></html>".encode(encoding)


class CatalogImportDisguisedFormatTestCase(TestCase):
    """Формат определяется по содержимому файла, а не по расширению."""

    def setUp(self) -> None:
        _map(CatalogType.DIRECTION, {"Направление": "name", "Код": "external_code"})

    def _import_directions(self, name: str, content: bytes) -> tuple[int, int]:
        return catalog_import_service.import_file(CatalogType.DIRECTION, SimpleUploadedFile(name, content))[:2]

    def test_html_table_saved_as_xls_is_read(self) -> None:
        content = _html_xls([("Направление", "Код"), ("Ростелеком", "rt")], charset="utf-8", encoding="utf-8")

        created, _ = self._import_directions("directions.xls", content)

        self.assertEqual(created, 1)
        self.assertTrue(Direction.objects.filter(name="Ростелеком", external_code="rt").exists())

    def test_html_in_cp1251_by_meta_charset(self) -> None:
        content = _html_xls([("Направление", "Код"), ("Ростелеком", "rt")], charset="windows-1251", encoding="cp1251")

        self._import_directions("directions.xls", content)

        self.assertTrue(Direction.objects.filter(name="Ростелеком").exists())

    def test_html_in_cp1251_without_charset_falls_back_to_cp1251(self) -> None:
        content = _html_xls([("Направление", "Код"), ("Ростелеком", "rt")], charset=None, encoding="cp1251")

        self._import_directions("directions.xls", content)

        self.assertTrue(Direction.objects.filter(name="Ростелеком").exists())

    def test_html_date_text_is_read_as_date(self) -> None:
        OrganizationFactory(name="МГУ")
        ProductFactory(name="1С:Предприятие", vendor=VendorFactory(name="1С"))
        columns = {
            "Вуз": "organization",
            "Вендор": "vendor",
            "Продукт": "product",
            "№ договора": "contract_number",
            "Лицензия подписана": "license_signed",
        }
        _map(CatalogType.CONTRACT_REGISTRY, columns)
        content = _html_xls(
            [tuple(columns), ("МГУ", "1С", "1С:Предприятие", "Д-1", "01.03.2026")], charset="utf-8", encoding="utf-8"
        )

        catalog_import_service.import_file(CatalogType.CONTRACT_REGISTRY, SimpleUploadedFile("registry.xls", content))

        self.assertEqual(License.objects.get(contract__contract_number="Д-1").signed_at, date(2026, 3, 1))

    def test_html_with_unclosed_cells_and_rows_is_read(self) -> None:
        content = (
            "<html><body><table>"
            "<tr><th>Направление<th>Код"
            "<tr><td>Ростелеком<td>rt"
            "<tr><td>JetBrains<td>jb"
            "</table></body></html>"
        ).encode("utf-8")

        created, _ = self._import_directions("directions.xls", content)

        self.assertEqual(created, 2)
        self.assertTrue(Direction.objects.filter(name="JetBrains", external_code="jb").exists())

    def test_text_saved_as_xls_raises_clear_error(self) -> None:
        content = "Направление\tКод\nРостелеком\trt\n".encode("cp1251")

        with self.assertRaises(CatalogImportError) as context:
            self._import_directions("directions.xls", content)

        self.assertIn("текст", str(context.exception))

    def test_real_xls_with_xlsx_extension_is_read(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xls"
            _write_xls(source, ("Направление", "Код"), ("Ростелеком", "rt"))
            content = source.read_bytes()

        created, _ = self._import_directions("directions.xlsx", content)

        self.assertEqual(created, 1)

    def test_old_xls_without_codepage_is_reread_as_cp1251(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xls"
            _write_xls(source, ("Направление", "Код"), ("Ростелеком", "rt"))
            content = source.read_bytes()
        real_open_workbook = xlrd.open_workbook
        calls = []

        def _open_workbook(**kwargs):
            calls.append(kwargs.get("encoding_override"))
            book = real_open_workbook(**kwargs)
            if kwargs.get("encoding_override") is None:
                # Имитируем файл Excel 95 без записи CODEPAGE.
                book.biff_version, book.codepage = 70, None
            return book

        with patch("sova.catalog.services.import_file.xlrd.open_workbook", side_effect=_open_workbook):
            self._import_directions("directions.xls", content)

        self.assertEqual(calls, [None, "cp1251"])


class CatalogImportTextNormalizationTestCase(TestCase):
    """Невидимые различия в тексте ячеек не ломают маппинг колонок и поиск по справочникам."""

    def test_header_matches_mapping_despite_case_nbsp_and_double_spaces(self) -> None:
        _map(CatalogType.DIRECTION, {"Наименование направления": "name", "Код": "external_code"})
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xlsx"
            _write_workbook(source, ("наименование  направления ", "КОД"), ("JetBrains", "jb"))

            created, _, _ = catalog_import_service.import_file(CatalogType.DIRECTION, source)

        self.assertEqual(created, 1)

    def test_cell_values_are_normalized(self) -> None:
        _map(CatalogType.DIRECTION, {"Направление": "name", "Код": "external_code"})
        decomposed = "Сервис Яндекс Облако \u0438\u0306"  # «й» в разложенном виде: «и» + кратка
        with TemporaryDirectory() as directory:
            source = Path(directory) / "directions.xlsx"
            _write_workbook(source, ("Направление", "Код"), (f"{decomposed} ​", "ya​cloud"))

            catalog_import_service.import_file(CatalogType.DIRECTION, source)

        direction = Direction.objects.get()
        self.assertEqual(direction.name, "Сервис Яндекс Облако й")
        self.assertEqual(direction.external_code, "yacloud")


class ReadHeadersTestCase(TestCase):
    """import_file_service.read_headers — заголовки файла для настройки маппинга."""

    def test_read_headers_returns_normalized_non_empty_headers_in_file_order(self) -> None:
        """Заголовки первой строки: нормализованы, пустые и повторы отброшены, порядок файла сохранён."""
        workbook = Workbook()
        workbook.active.append(("Вендор ", None, "Код\u00a0товара", "Вендор"))
        workbook.active.append(("JetBrains", "x", "jb", "y"))
        content = BytesIO()
        workbook.save(content)

        headers = import_file_service.read_headers(source=SimpleUploadedFile("v.xlsx", content.getvalue()))

        # Проверяем состав и порядок заголовков
        self.assertEqual(headers, ["Вендор", "Код товара"])

    def test_read_headers_joins_multiline_header_into_one_line(self) -> None:
        """Перенос строки в заголовке — пробел: так заголовок можно выбрать в однострочном поле маппинга."""
        workbook = Workbook()
        workbook.active.append(("Дата\nподписания",))
        content = BytesIO()
        workbook.save(content)

        headers = import_file_service.read_headers(source=SimpleUploadedFile("r.xlsx", content.getvalue()))

        # Проверяем, что заголовок отдан в одну строку
        self.assertEqual(headers, ["Дата подписания"])

    def test_read_headers_of_empty_file_raises(self) -> None:
        """Пустой файл — ошибка файла, как при импорте."""
        content = BytesIO()
        Workbook().save(content)

        # Проверяем, что пустой файл отклоняется
        with self.assertRaises(CatalogImportError):
            import_file_service.read_headers(source=SimpleUploadedFile("empty.xlsx", content.getvalue()))
