from collections.abc import Iterator
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import CommandError
from django.core.management.commands.loaddata import Command as DjangoLoadDataCommand
from django.db import transaction
from openpyxl import load_workbook

from sova.catalog.models import Direction, Product, Program, University, Vendor
from sova.catalog.models.university import InstitutionType


class Command(DjangoLoadDataCommand):
    help = "Загружает Django fixtures или справочники XLSX из reference_data."

    def add_arguments(self, parser) -> None:
        super().add_arguments(parser)
        # Django объявляет метки фикстур обязательными (nargs="+"), но с --reference-data
        # указывать их не нужно: ослабляем до nargs="*" и проверяем комбинацию сами.
        for action in parser._actions:
            if action.dest == "args":
                action.nargs = "*"
                break
        parser.add_argument(
            "--reference-data",
            action="store_true",
            help=(
                "Загрузить все справочники XLSX из reference_data "
                "в порядке зависимостей, без указания файлов."
            ),
        )

    def handle(self, *fixture_labels, **options):
        if options.pop("reference_data", False):
            if fixture_labels:
                raise CommandError("--reference-data нельзя совмещать с указанием файлов.")
            return self._load_reference_data()

        if not fixture_labels:
            raise CommandError(self.missing_args_message)

        source = self._resolve_xlsx(fixture_labels)
        if source is None:
            return super().handle(*fixture_labels, **options)

        loader = self._loaders().get(source.name)
        if loader is None:
            raise CommandError(f"Неизвестный XLSX-справочник: {source.name}")
        loader(source)

    def _loaders(self) -> dict:
        """Справочники в порядке зависимостей: продукты ссылаются на вендоров и программы."""
        return {
            "universities.xlsx": self._load_universities,
            "vendors.xlsx": self._load_vendors,
            "directions.xlsx": self._load_directions,
            "programs.xlsx": self._load_programs,
            "products.xlsx": self._load_products,
        }

    @transaction.atomic
    def _load_reference_data(self) -> None:
        directory = Path(settings.BASE_DIR) / "reference_data"
        loaders = self._loaders()

        missing = [name for name in loaders if not (directory / name).is_file()]
        if missing:
            raise CommandError(f"В {directory} не найдены справочники: {', '.join(missing)}")

        for name, loader in loaders.items():
            loader(directory / name)

        self.stdout.write(self.style.SUCCESS(f"Загружены все справочники: {', '.join(loaders)}."))

    def _resolve_xlsx(self, fixture_labels: tuple[str, ...]) -> Path | None:
        if len(fixture_labels) != 1:
            return None

        label = Path(fixture_labels[0])
        if label.suffix.lower() != ".xlsx":
            return None

        candidates = [label]
        if not label.is_absolute():
            candidates.append(Path(settings.BASE_DIR) / "reference_data" / label.name)

        for candidate in candidates:
            if candidate.is_file():
                return candidate

        raise CommandError(f"Файл справочника не найден: {fixture_labels[0]}")

    def _read_rows(self, source: Path, required: set[str]) -> Iterator[tuple[int, dict]]:
        workbook = load_workbook(source, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)

        try:
            raw_headers = next(rows)
        except StopIteration as error:
            workbook.close()
            raise CommandError(f"Файл {source.name} пуст.") from error

        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        missing = sorted(required - set(headers))
        if missing:
            workbook.close()
            raise CommandError(f"В {source.name} отсутствуют колонки: {', '.join(missing)}")

        for row_number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            yield row_number, dict(zip(headers, values, strict=False))

        workbook.close()

    @transaction.atomic
    def _load_universities(self, source: Path) -> None:
        required = {
            "id",
            "ror",
            "name_en",
            "name",
            "short_name",
            "country_code",
            "type",
            "works_count",
            "cited_by_count",
            "city",
            "region",
            "lat",
            "lon",
            "homepage_url",
        }
        created = 0
        updated = 0
        allowed_types = {value for value, _label in InstitutionType.choices}

        for row_number, row in self._read_rows(source, required):
            try:
                name = self._text(row["name"])
                if not name:
                    raise ValueError("поле name обязательно")

                external_code = self._text(row["ror"]) or self._text(row["id"])
                if not external_code:
                    raise ValueError("должно быть заполнено поле ror или id")

                institution_type = self._text(row["type"]).lower() or InstitutionType.EDUCATION
                if institution_type not in allowed_types:
                    raise ValueError(f"неизвестный type: {institution_type}")

                defaults = {
                    "name": name,
                    "name_en": self._text(row["name_en"]),
                    "short_name": self._text(row["short_name"]),
                    "institution_type": institution_type,
                    "country_code": self._text(row["country_code"]).upper(),
                    "region": self._text(row["region"]),
                    "city": self._text(row["city"]),
                    "lat": self._decimal(row["lat"]),
                    "lon": self._decimal(row["lon"]),
                    "works_count": self._positive_int(row["works_count"]),
                    "cited_by_count": self._positive_int(row["cited_by_count"]),
                    "homepage_url": self._text(row["homepage_url"]),
                }
            except (InvalidOperation, TypeError, ValueError) as error:
                raise CommandError(f"Ошибка в строке {row_number}: {error}") from error

            _, was_created = self._upsert(University, external_code, name, defaults)
            created += was_created
            updated += not was_created

        self.stdout.write(
            self.style.SUCCESS(f"Справочник вузов загружен: создано {created}, обновлено {updated}.")
        )

    @transaction.atomic
    def _load_vendors(self, source: Path) -> None:
        required = {"name", "external_code"}
        created = 0
        updated = 0

        for row_number, row in self._read_rows(source, required):
            try:
                name = self._text(row["name"])
                if not name:
                    raise ValueError("поле name обязательно")
                external_code = self._text(row["external_code"]) or None
                is_active = self._bool(row.get("is_active"))
            except ValueError as error:
                raise CommandError(f"Ошибка в строке {row_number}: {error}") from error

            _, was_created = self._upsert(Vendor, external_code, name, {"name": name, "is_active": is_active})
            created += was_created
            updated += not was_created

        self.stdout.write(
            self.style.SUCCESS(f"Справочник вендоров загружен: создано {created}, обновлено {updated}.")
        )

    @transaction.atomic
    def _load_directions(self, source: Path) -> None:
        required = {"name", "external_code"}
        created = 0
        updated = 0

        for row_number, row in self._read_rows(source, required):
            try:
                name = self._text(row["name"])
                if not name:
                    raise ValueError("поле name обязательно")
                external_code = self._text(row["external_code"]) or None
                is_active = self._bool(row.get("is_active"))
            except ValueError as error:
                raise CommandError(f"Ошибка в строке {row_number}: {error}") from error

            _, was_created = self._upsert(
                Direction, external_code, name, {"name": name, "is_active": is_active}
            )
            created += was_created
            updated += not was_created

        self.stdout.write(
            self.style.SUCCESS(f"Справочник направлений загружен: создано {created}, обновлено {updated}.")
        )

    @transaction.atomic
    def _load_products(self, source: Path) -> None:
        required = {"name", "external_code", "vendor"}
        created = 0
        updated = 0

        for row_number, row in self._read_rows(source, required):
            try:
                name = self._text(row["name"])
                if not name:
                    raise ValueError("поле name обязательно")
                external_code = self._text(row["external_code"]) or None
                is_active = self._bool(row.get("is_active"))
            except ValueError as error:
                raise CommandError(f"Ошибка в строке {row_number}: {error}") from error

            vendor = self._find_vendor(row_number, row["vendor"])
            defaults = {"name": name, "vendor": vendor, "is_active": is_active}

            # Уникальность продукта — в паре с вендором (или одна, если вендора нет), поэтому
            # апсерт по имени ищет среди продуктов того же вендора, а не по всему справочнику.
            product = None
            if external_code:
                product = Product.objects.filter(external_code=external_code).first()
            if product is None:
                product = Product.objects.filter(name=name, vendor=vendor).first()

            if product is None:
                product = Product.objects.create(external_code=external_code, **defaults)
                created += 1
            else:
                product.external_code = external_code
                for field, value in defaults.items():
                    setattr(product, field, value)
                product.save(update_fields=["external_code", *defaults, "updated_at"])
                updated += 1

            # Колонка programs необязательна: если её нет в файле — существующие связи не трогаем,
            # если есть (пусть и пустая) — приводим M2M к тому, что в ней перечислено.
            if "programs" in row:
                product.programs.set(self._find_programs(row_number, row["programs"]))

        self.stdout.write(
            self.style.SUCCESS(f"Справочник продуктов загружен: создано {created}, обновлено {updated}.")
        )

    @transaction.atomic
    def _load_programs(self, source: Path) -> None:
        required = {"name", "direction"}
        created = 0
        updated = 0

        for row_number, row in self._read_rows(source, required):
            try:
                name = self._text(row["name"])
                if not name:
                    raise ValueError("поле name обязательно")
                is_active = self._bool(row.get("is_active"))
            except ValueError as error:
                raise CommandError(f"Ошибка в строке {row_number}: {error}") from error

            direction = self._find_direction(row_number, row["direction"])
            defaults = {"name": name, "direction": direction, "is_active": is_active}

            # У Program нет external_code, поэтому апсерт идёт по паре name+direction
            # (аналогично name+vendor у Product).
            program = Program.objects.filter(name=name, direction=direction).first()
            if program is None:
                Program.objects.create(**defaults)
                created += 1
            else:
                for field, value in defaults.items():
                    setattr(program, field, value)
                program.save(update_fields=[*defaults, "updated_at"])
                updated += 1

        self.stdout.write(
            self.style.SUCCESS(f"Справочник программ загружен: создано {created}, обновлено {updated}.")
        )

    def _find_vendor(self, row_number: int, raw_value) -> Vendor | None:
        value = self._text(raw_value)
        if not value:
            return None
        vendor = Vendor.objects.filter(external_code=value).first() or Vendor.objects.filter(name=value).first()
        if vendor is None:
            raise CommandError(f"Ошибка в строке {row_number}: вендор не найден: {value}")
        return vendor

    def _find_direction(self, row_number: int, raw_value) -> Direction:
        value = self._text(raw_value)
        if not value:
            raise CommandError(f"Ошибка в строке {row_number}: поле direction обязательно")
        direction = (
            Direction.objects.filter(external_code=value).first()
            or Direction.objects.filter(name=value).first()
        )
        if direction is None:
            raise CommandError(f"Ошибка в строке {row_number}: направление не найдено: {value}")
        return direction

    def _find_programs(self, row_number: int, raw_value) -> list[Program]:
        value = self._text(raw_value)
        if not value:
            return []

        programs = []
        for name in (part.strip() for part in value.split(";")):
            if not name:
                continue
            program = Program.objects.filter(name=name).first()
            if program is None:
                raise CommandError(f"Ошибка в строке {row_number}: программа не найдена: {name}")
            programs.append(program)
        return programs

    @staticmethod
    def _upsert(model, external_code: str | None, name: str, defaults: dict) -> tuple[object, bool]:
        """Апсерт справочника по external_code (если задан), иначе по name."""
        instance = None
        if external_code:
            instance = model.objects.filter(external_code=external_code).first()
        if instance is None:
            instance = model.objects.filter(name=name).first()

        if instance is None:
            return model.objects.create(external_code=external_code, **defaults), True

        instance.external_code = external_code
        for field, value in defaults.items():
            setattr(instance, field, value)
        instance.save(update_fields=["external_code", *defaults, "updated_at"])
        return instance, False

    @staticmethod
    def _text(value) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _bool(value) -> bool:
        if value is None or value == "":
            return True
        if isinstance(value, bool):
            return value
        text = str(value).strip().lower()
        if text in {"1", "true", "да", "yes"}:
            return True
        if text in {"0", "false", "нет", "no"}:
            return False
        raise ValueError(f"ожидалось булево значение is_active, получено {value}")

    @staticmethod
    def _decimal(value) -> Decimal | None:
        if value in (None, ""):
            return None
        return Decimal(str(value))

    @staticmethod
    def _positive_int(value) -> int:
        if value in (None, ""):
            return 0
        number = Decimal(str(value))
        if number < 0 or number != number.to_integral_value():
            raise ValueError(f"ожидалось целое неотрицательное число, получено {value}")
        return int(number)
